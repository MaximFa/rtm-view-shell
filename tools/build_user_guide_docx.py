"""Convert USER_GUIDE.md to a formatted user-instruction DOCX."""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT     = Path(__file__).parent.parent
MD_FILE  = ROOT / "USER_GUIDE.md"
OUT_FILE = ROOT / "RTM_View_Shell_User_Guide.docx"

# ── Palette ─────────────────────────────────────────────────────────
BLUE_DARK  = RGBColor(0x1F, 0x3A, 0x5F)
BLUE_MID   = RGBColor(0x2E, 0x75, 0xB6)
BLUE_LIGHT = RGBColor(0xBD, 0xD7, 0xEE)
GREEN      = RGBColor(0x37, 0x86, 0x44)
AMBER      = RGBColor(0x9C, 0x6A, 0x00)
GREY_TEXT  = RGBColor(0x40, 0x40, 0x40)
NOTE_BG    = "EBF3FB"   # light blue for blockquotes
CODE_BG    = "F2F2F2"


# ── XML helpers ─────────────────────────────────────────────────────

def set_para_shading(para, fill_hex: str):
    pPr = para._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def set_cell_shading(cell, fill_hex: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tcPr.append(shd)


def add_rule(doc: Document, color: str = "CCCCCC"):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "4")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), color)
    pb.append(bot)
    pPr.append(pb)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after  = Pt(4)


# ── Style configuration ──────────────────────────────────────────────

def setup_styles(doc: Document):
    # Normal
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(11)
    n.font.color.rgb = GREY_TEXT

    # Headings
    specs = [
        ("Heading 1", 20, BLUE_DARK, True,  18, 6),
        ("Heading 2", 14, BLUE_MID,  True,  14, 4),
        ("Heading 3", 11, BLUE_DARK, True,  10, 3),
    ]
    for name, sz, color, bold, before, after in specs:
        s = doc.styles[name]
        s.font.name  = "Calibri"
        s.font.size  = Pt(sz)
        s.font.color.rgb = color
        s.font.bold  = bold
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after  = Pt(after)
        s.paragraph_format.keep_with_next = True

    # Code Block
    if "Code Block" not in [s.name for s in doc.styles]:
        cs = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cs = doc.styles["Code Block"]
    cs.font.name = "Courier New"
    cs.font.size = Pt(9)
    cs.paragraph_format.left_indent  = Cm(0.6)
    cs.paragraph_format.right_indent = Cm(0.6)
    cs.paragraph_format.space_before = Pt(3)
    cs.paragraph_format.space_after  = Pt(3)

    # Inline Code
    if "ICode" not in [s.name for s in doc.styles]:
        ic = doc.styles.add_style("ICode", WD_STYLE_TYPE.CHARACTER)
    else:
        ic = doc.styles["ICode"]
    ic.font.name = "Courier New"
    ic.font.size = Pt(10)
    ic.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)

    # Note (blockquote)
    if "Note" not in [s.name for s in doc.styles]:
        note = doc.styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        note = doc.styles["Note"]
    note.font.name   = "Calibri"
    note.font.size   = Pt(10.5)
    note.font.italic = True
    note.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    note.paragraph_format.left_indent  = Cm(0.5)
    note.paragraph_format.right_indent = Cm(0.5)
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after  = Pt(4)


# ── Title page ───────────────────────────────────────────────────────

def title_page(doc: Document):
    doc.add_paragraph()
    doc.add_paragraph()

    # Logo-like block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RTM VIEW SHELL")
    r.font.name  = "Calibri"
    r.font.size  = Pt(32)
    r.font.bold  = True
    r.font.color.rgb = BLUE_DARK

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("User Guide")
    r2.font.name  = "Calibri"
    r2.font.size  = Pt(20)
    r2.font.color.rgb = BLUE_MID

    doc.add_paragraph()

    # Decorative rule
    add_rule(doc, "2E75B6")

    # Subtitle
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Real-Time Monitoring Dashboard Platform")
    sr.font.name  = "Calibri"
    sr.font.size  = Pt(12)
    sr.font.color.rgb = GREY_TEXT
    sr.font.italic = True

    add_rule(doc, "2E75B6")
    doc.add_page_break()


# ── Table of contents (static) ───────────────────────────────────────

def toc_page(doc: Document, sections: list[tuple[str, str]]):
    doc.add_heading("Contents", level=1)
    for num, title in sections:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after  = Pt(3)
        r = p.add_run(f"{num}  {title}")
        r.font.size = Pt(11)
    doc.add_page_break()


# ── Inline formatting helpers ─────────────────────────────────────────

def para_with_markup(doc: Document, text: str, style: str = "Normal"):
    """Add paragraph, converting `code`, **bold**, and plain text."""
    para = doc.add_paragraph(style=style)
    _fill_runs(para, text)
    return para


def _fill_runs(para, text: str):
    # Strip markdown links
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    # Split on backticks and **bold**
    tokens = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for tok in tokens:
        if tok.startswith("`") and tok.endswith("`"):
            r = para.add_run(tok[1:-1])
            r.style = "ICode"
        elif tok.startswith("**") and tok.endswith("**"):
            r = para.add_run(tok[2:-2])
            r.bold = True
        else:
            para.add_run(tok)


# ── Block builders ───────────────────────────────────────────────────

def add_code_block(doc: Document, lines: list[str]):
    p = doc.add_paragraph(style="Code Block")
    set_para_shading(p, CODE_BG)
    first = True
    for line in lines:
        if first:
            p.add_run(line)
            first = False
        else:
            p.add_run("\n" + line)
    p.paragraph_format.keep_together = True


def add_note(doc: Document, text: str):
    """Render a blockquote (> …) as a shaded note box."""
    p = doc.add_paragraph(style="Note")
    set_para_shading(p, NOTE_BG)
    # icon
    icon_run = p.add_run("ℹ  ")
    icon_run.font.bold  = True
    icon_run.font.color.rgb = BLUE_MID
    icon_run.font.italic = False
    _fill_runs(p, text)


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    tbl  = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = "Table Grid"
    for i, row_data in enumerate(rows):
        for j in range(cols):
            cell = tbl.cell(i, j)
            text = row_data[j] if j < len(row_data) else ""
            text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)   # strip bold markers
            text = re.sub(r"`([^`]+)`", r"\1", text)           # strip code markers
            cell.text = text
            if i == 0:
                set_cell_shading(cell, "1F3A5F")
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(10.5)
            else:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(10.5)
    doc.add_paragraph()


def add_step_list(doc: Document, items: list[str], ordered: bool):
    for idx, text in enumerate(items, 1):
        style = "List Number" if ordered else "List Bullet"
        p = para_with_markup(doc, text, style=style)
        p.paragraph_format.left_indent       = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        p.paragraph_format.space_before      = Pt(2)
        p.paragraph_format.space_after       = Pt(2)


# ── Main parser ──────────────────────────────────────────────────────

def render_markdown(doc: Document, md: str):
    lines = md.split("\n")
    i = 0

    code_buf: list[str] = []
    in_code = False
    table_buf: list[str] = []
    in_table = False
    list_buf: list[str] = []
    list_ordered = False
    in_list = False

    def flush_list():
        nonlocal list_buf, in_list
        if list_buf:
            add_step_list(doc, list_buf, list_ordered)
        list_buf = []
        in_list  = False

    def flush_table():
        nonlocal table_buf, in_table
        if table_buf:
            rows = []
            for ln in table_buf:
                if re.match(r"\s*\|[-:| ]+\|\s*$", ln):
                    continue
                cells = [c.strip() for c in ln.strip().strip("|").split("|")]
                rows.append(cells)
            add_table(doc, rows)
        table_buf = []
        in_table  = False

    while i < len(lines):
        line = lines[i]

        # ── Code fence ────────────────────────────────────────────
        if line.strip().startswith("```"):
            flush_list(); flush_table()
            if not in_code:
                in_code = True
                code_buf = []
            else:
                add_code_block(doc, code_buf)
                in_code = False
                code_buf = []
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # ── Table ─────────────────────────────────────────────────
        if line.startswith("|"):
            flush_list()
            in_table = True
            table_buf.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()

        # ── Blank line ─────────────────────────────────────────────
        if not line.strip():
            flush_list()
            i += 1
            continue

        # ── HR ─────────────────────────────────────────────────────
        if re.match(r"^---+\s*$", line):
            flush_list()
            add_rule(doc)
            i += 1
            continue

        # ── Headings ───────────────────────────────────────────────
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            flush_list()
            level = len(m.group(1))
            text  = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", m.group(2))
            doc.add_heading(text, level=level)
            i += 1
            continue

        # ── Blockquote ─────────────────────────────────────────────
        if line.startswith("> "):
            flush_list()
            add_note(doc, line[2:])
            i += 1
            continue

        # ── Ordered list ───────────────────────────────────────────
        m = re.match(r"^\d+\.\s+(.*)", line)
        if m:
            if in_list and not list_ordered:
                flush_list()
            in_list = True
            list_ordered = True
            list_buf.append(m.group(1))
            i += 1
            continue

        # ── Unordered list ─────────────────────────────────────────
        m = re.match(r"^[-*]\s+(.*)", line)
        if m:
            if in_list and list_ordered:
                flush_list()
            in_list = True
            list_ordered = False
            list_buf.append(m.group(1))
            i += 1
            continue

        # ── Normal paragraph ───────────────────────────────────────
        flush_list()
        para_with_markup(doc, line)
        i += 1

    flush_list()
    flush_table()


# ── Entry point ──────────────────────────────────────────────────────

TOC = [
    ("1", "Overview"),
    ("2", "Signing In"),
    ("  2.1", "Email and Password"),
    ("  2.2", "Single Sign-On (SSO)"),
    ("  2.3", "Two-Factor Authentication (2FA)"),
    ("3", "Application Layout"),
    ("4", "Dashboards"),
    ("  4.1", "Viewing Your Dashboards"),
    ("  4.2", "Creating a Dashboard"),
    ("  4.3", "Editing a Dashboard"),
    ("  4.4", "Deleting a Dashboard"),
    ("5", "Widget Catalogue"),
    ("6", "Administration"),
    ("  6.1", "User Management"),
    ("  6.2", "Permission Groups"),
    ("7", "Signing Out"),
    ("8", "Frequently Asked Questions"),
]


def main():
    md = MD_FILE.read_text(encoding="utf-8")

    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.0)

    setup_styles(doc)
    title_page(doc)
    toc_page(doc, TOC)

    # Drop the H1 title line — already on title page
    lines = md.split("\n")
    body_start = next((i for i, l in enumerate(lines) if l.startswith("## ")), 0)
    render_markdown(doc, "\n".join(lines[body_start:]))

    doc.save(OUT_FILE)
    print(f"Saved: {OUT_FILE}")


if __name__ == "__main__":
    main()
